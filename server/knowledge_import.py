"""A file dropped on «Знання» → records for a knowledge set (server/knowledge.py), shown to the user before
anything is saved.

CSV, Excel and JSON are tables already: a row is a record, the name column its title, the other columns
its facts. DOCX tables are read the same way; free text (DOCX paragraphs, TXT, MD) goes to the backend
model, which pulls out one record per person (or per topic, when the text is not about people).
"""
from __future__ import annotations

import csv
import io
import json
import logging
import re
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from server.knowledge import MAX_ITEMS_PER_SET, clean_items

logger = logging.getLogger(__name__)

MAX_FILE_BYTES = 5_000_000
MAX_TEXT_FOR_MODEL = 60_000
EXTENSIONS = ("csv", "tsv", "xlsx", "json", "docx", "txt", "md")

Extract = Callable[[str], list[dict[str, Any]]]
# header + a few sample rows → {"title": [column index…], "labels": {index: short label}, "skip": [index…]}
Plan = Callable[[list[str], list[list[str]]], dict[str, Any]]
MAX_LABEL_CHARS = 60
PLAN_SAMPLE_ROWS = 5
PLAN_SAMPLE_CHARS = 80


class KnowledgeImportError(Exception):
    """The file can't be read; the message is for the page."""


@dataclass
class Parsed:
    items: list[dict[str, Any]]
    kind: str  # "table" | "text" | "mixed"
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        columns: list[str] = []
        for item in self.items:
            for key in item["fields"]:
                if key not in columns:
                    columns.append(key)
        return {"items": self.items, "columns": columns, "kind": self.kind, "warnings": self.warnings}


# Which column names a person (or the record): a full name first, else first + last name, else the first column.
_FULL_NAME = re.compile(
    r"^(піб|п\.?\s?і\.?\s?б\.?|фіо|фио|ф\.?\s?и\.?\s?о\.?|full\s*name|name|ім.я\s*(та|і|й)\s*прізвище|"
    r"прізвище\s*(та|і|й)?\s*ім.я|имя\s*(и)?\s*фамилия|фамилия\s*(и)?\s*имя|хто|людина|person|співробітник|"
    r"сотрудник|учасник|участник|гість|гость|title|назва|название)$", re.I)
# A form's own wording («Твоє Ім'я та прізвище», «Your full name»): the same meaning, anywhere in the header.
_FULL_NAME_LOOSE = re.compile(
    r"(ім.я.{0,12}прізвищ|прізвищ.{0,12}ім.я|піб\b|фіо\b|имя.{0,12}фамили|фамили.{0,12}имя|full\s*name|your\s*name)", re.I)
# Form exports start with the time of the answer: never a person's name.
_TIME_COLUMN = re.compile(r"(позначка\s*часу|відмітка\s*часу|мітка\s*часу|отметка\s*времени|timestamp|submitted)", re.I)
_FIRST_NAME = re.compile(r"^(ім.я|имя|first\s*name|given\s*name|firstname)$", re.I)
_LAST_NAME = re.compile(r"^(прізвище|фамилия|last\s*name|surname|lastname|family\s*name)$", re.I)


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(" ", "minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, list):
        return ", ".join(_cell(v) for v in value if _cell(v))
    if isinstance(value, dict):
        return "; ".join(f"{k}: {_cell(v)}" for k, v in value.items() if _cell(v))
    return " ".join(str(value).split())


def _planned(plan: Plan | None, header: list[str], rows: list[list[str]]) -> dict[str, Any]:
    """The model's view of the table (which column names the person, short labels for long question headers,
    technical columns to leave out), checked: anything off and that part is left to the plain rules."""
    if plan is None:
        return {}
    width = len(header)
    sample = [[c[:PLAN_SAMPLE_CHARS] for c in r[:width]] for r in rows[1:1 + PLAN_SAMPLE_ROWS]]
    try:
        said = plan(header, sample)
    except Exception:
        logger.exception("knowledge.plan_failed")
        return {}
    if not isinstance(said, dict):
        return {}

    def indexes(value: Any) -> list[int]:
        return sorted({int(i) for i in value if isinstance(i, int) and not isinstance(i, bool) and 0 <= i < width}) \
            if isinstance(value, list) else []

    labels: dict[int, str] = {}
    raw = said.get("labels")
    for key, value in (raw.items() if isinstance(raw, dict) else []):
        label = " ".join(str(value or "").replace("?", " ").split()).strip(" :.,;")[:MAX_LABEL_CHARS]
        if str(key).isdigit() and 0 <= int(key) < width and label:
            labels[int(key)] = label
    title = indexes(said.get("title"))
    skip = [i for i in indexes(said.get("skip")) if i not in title]
    return {"title": title[:2], "labels": labels, "skip": skip}


def table_to_items(rows: list[list[Any]], plan: Plan | None = None) -> list[dict[str, Any]]:
    """Header row + rows → records. Empty rows are skipped; the name column becomes the title. With a plan
    (the model's), it picks the title column and short labels; the plain rules cover whatever it left out."""
    rows = [[_cell(c) for c in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if len(rows) < 2:
        return []
    width = max(len(r) for r in rows)
    header: list[str] = []
    for i in range(width):
        name = rows[0][i] if i < len(rows[0]) else ""
        name = name or f"Колонка {i + 1}"
        while name in header:
            name += " (2)"
        header.append(name)
    full = next((i for i, h in enumerate(header) if _FULL_NAME.match(h.strip())), None)
    first = next((i for i, h in enumerate(header) if _FIRST_NAME.match(h.strip())), None)
    last = next((i for i, h in enumerate(header) if _LAST_NAME.match(h.strip())), None)
    if full is None:
        full = next((i for i, h in enumerate(header) if _FULL_NAME_LOOSE.search(h)), None)
    planned = _planned(plan, header, rows)
    if planned.get("title"):
        title_cols = planned["title"]
    elif full is not None:
        title_cols = [full]
    elif first is not None or last is not None:
        title_cols = [i for i in (first, last) if i is not None]
    else:
        title_cols = [next((i for i, h in enumerate(header) if not _TIME_COLUMN.search(h)), 0)]
    skip = set(planned.get("skip", []))
    labels: list[str] = []
    for i, name in enumerate(header):
        label = planned.get("labels", {}).get(i, name)
        while label in labels:
            label += " (2)"
        labels.append(label)
    items = []
    for row in rows[1:]:
        row = row + [""] * (width - len(row))
        title = " ".join(row[i] for i in title_cols if row[i])
        fields = {labels[i]: row[i] for i in range(width) if i not in title_cols and i not in skip and row[i]}
        items.append({"title": title, "fields": fields})
    return items


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _csv(data: bytes, ext: str, plan: Plan | None = None) -> list[dict[str, Any]]:
    text = _decode(data)
    if ext == "tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    return table_to_items(list(csv.reader(io.StringIO(text), delimiter=delimiter)), plan)


def _xlsx(data: bytes, plan: Plan | None = None) -> tuple[list[dict[str, Any]], list[str]]:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise KnowledgeImportError("Не вдалося відкрити Excel-файл (потрібен .xlsx).") from exc
    items: list[dict[str, Any]] = []
    sheets = []
    for sheet in book.worksheets:
        found = table_to_items([list(r) for r in sheet.iter_rows(values_only=True)], plan)
        if found:
            items += found
            sheets.append(sheet.title)
    book.close()
    warnings = [f"Узято аркуші: {', '.join(sheets)}."] if len(sheets) > 1 else []
    return items, warnings


def _json(data: bytes, plan: Plan | None = None) -> list[dict[str, Any]]:
    try:
        doc = json.loads(_decode(data))
    except json.JSONDecodeError as exc:
        raise KnowledgeImportError(f"JSON не читається: рядок {exc.lineno}, {exc.msg}.") from exc
    if isinstance(doc, dict):
        lists = [v for v in doc.values() if isinstance(v, list) and v and all(isinstance(x, dict) for x in v)]
        if lists:
            doc = lists[0]  # {"people": [...]}
        elif all(isinstance(v, dict) for v in doc.values()):
            doc = [{"__title": k, **v} for k, v in doc.items()]  # {"Олег": {...}, "Марина": {...}}
        else:
            doc = [doc]
    if not isinstance(doc, list):
        raise KnowledgeImportError("У JSON очікується список записів або об'єкт зі списком.")
    rows: list[dict[str, Any]] = [x if isinstance(x, dict) else {"__title": x} for x in doc]
    keys: list[str] = []
    for row in rows:
        for key in row:
            if key not in keys:
                keys.append(key)
    keys.sort(key=lambda k: k != "__title")
    table = [["ПІБ" if k == "__title" else k for k in keys]] + [[row.get(k) for k in keys] for row in rows]
    return table_to_items(table, plan)


def _docx(data: bytes, plan: Plan | None = None) -> tuple[list[dict[str, Any]], str]:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise KnowledgeImportError("Не вдалося відкрити DOCX (старий .doc збережіть як .docx).") from exc
    items: list[dict[str, Any]] = []
    for table in doc.tables:
        items += table_to_items([[cell.text for cell in row.cells] for row in table.rows], plan)
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return items, text


def parse_file(filename: str, data: bytes, extract: Extract | None = None, plan: Plan | None = None) -> Parsed:
    if not data:
        raise KnowledgeImportError("Файл порожній.")
    if len(data) > MAX_FILE_BYTES:
        raise KnowledgeImportError(f"Файл завеликий: до {MAX_FILE_BYTES // 1_000_000} МБ.")
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "doc" or ext == "xls":
        raise KnowledgeImportError(f"Старий формат .{ext}: збережіть його як .{ext}x і завантажте ще раз.")
    if ext not in EXTENSIONS:
        raise KnowledgeImportError("Підтримуються CSV, Excel (.xlsx), JSON, DOCX і текст (.txt, .md).")
    warnings: list[str] = []
    text = ""
    kind = "table"
    if ext in ("csv", "tsv"):
        items = _csv(data, ext, plan)
    elif ext == "xlsx":
        items, warnings = _xlsx(data, plan)
    elif ext == "json":
        items = _json(data, plan)
    elif ext == "docx":
        items, text = _docx(data, plan)
    else:
        items, text = [], _decode(data)
    # Free text worth reading: a document's paragraphs (beyond a title line) or a whole text file.
    if text.strip() and (not items or len(text.strip()) >= 200):
        if extract is None:
            raise KnowledgeImportError("Розбір тексту зараз недоступний — завантажте таблицю (CSV, Excel, JSON).")
        if len(text) > MAX_TEXT_FOR_MODEL:
            warnings.append(f"Текст довгий: розібрано перші {MAX_TEXT_FOR_MODEL // 1000} тис. символів.")
        try:
            found = extract(text[:MAX_TEXT_FOR_MODEL])
        except Exception as exc:
            logger.exception("knowledge.extract_failed")
            raise KnowledgeImportError("Не вдалося розібрати текст — спробуйте ще раз або завантажте таблицю.") from exc
        kind = "mixed" if items else "text"
        items += found
        warnings.append("Записи з тексту витягла модель — перевірте їх перед збереженням.")
    if len(items) > MAX_ITEMS_PER_SET:
        warnings.append(f"Записів {len(items)}: збережуться перші {MAX_ITEMS_PER_SET}.")
    items = clean_items(items)
    if not items:
        raise KnowledgeImportError("У файлі не знайшлося записів. Таблиця: перший рядок — назви колонок, далі — по рядку на людину.")
    return Parsed(items=items, kind=kind, warnings=warnings)


_EXTRACT_PROMPT = (
    "Extract records from this document for a voice assistant's knowledge base about the user's team. "
    "Each person is one record: title = their full name as written; fields = their facts as short "
    "key → value pairs in the document's language and wording (e.g. «вік»: «34», «хобі»: «риболовля, "
    "шахи», «улюблені фільми»: «Кримінальне чтиво», «посада»: «дизайнерка»). If the document is not about "
    "people (an event, a programme, rules), make one record per distinct topic: title = the topic, "
    "fields = its facts. Copy facts only — never invent, guess or add anything that is not in the text. "
    "Return JSON only, no prose: {\"items\": [{\"title\": \"...\", \"fields\": {\"...\": \"...\"}}]}\n\n"
    "Document:\n"
)


def make_openai_extract(client: Any, model: str) -> Extract:
    """Free text → records with the Responses API (sync client; runs in a worker thread)."""

    def extract(text: str) -> list[dict[str, Any]]:
        response = client.responses.create(model=model, input=_EXTRACT_PROMPT + text)
        out = (response.output_text or "").strip()
        out = re.sub(r"^```(?:json)?\s*|\s*```$", "", out)
        start, end = out.find("{"), out.rfind("}")
        doc = json.loads(out[start:end + 1]) if start >= 0 else {}
        items = doc.get("items") if isinstance(doc, dict) else None
        if not isinstance(items, list):
            raise ValueError("no items in the model's answer")
        return [
            {"title": _cell(x.get("title")), "fields": {str(k): _cell(v) for k, v in (x.get("fields") or {}).items()}}
            for x in items if isinstance(x, dict)
        ]

    return extract


_PLAN_PROMPT = (
    "A table is being added to a voice assistant's knowledge base about the user's team. Below: its column "
    "headers with their index, and a few sample rows. Return JSON only, no prose:\n"
    "{\"title\": [column indexes], \"labels\": {\"index\": \"short label\"}, \"skip\": [column indexes]}\n"
    "- title: the column (or two: first name + surname) holding the person's name — one record per row is a "
    "person. If the rows are not people, the column that names the row; if none, an empty list. Never a "
    "timestamp, date, number or id column.\n"
    "- labels: for every other column, a short noun-phrase label in the table's language, at most 5 words, "
    "no question mark, keeping the meaning: «Який твій улюблений фільм, серіал або гра?» → «Улюблений "
    "фільм, серіал, гра»; «Два цікаві факти про себе» → «Цікаві факти»; «Скільки тобі років?» → «Вік». "
    "A header that is already short and clear stays as it is.\n"
    "- skip: only purely technical columns that say nothing about the person — a form's submission timestamp, "
    "row numbers, ids. Never a column with answers, even a boring one.\n\n"
)


def make_openai_plan(client: Any, model: str) -> Plan:
    """Headers + sample rows → how to read the table, with the Responses API (sync; runs in a worker thread)."""

    def plan(header: list[str], sample: list[list[str]]) -> dict[str, Any]:
        table = "Headers:\n" + "\n".join(f"{i}: {h}" for i, h in enumerate(header))
        table += "\n\nSample rows:\n" + "\n".join(" | ".join(f"{i}={c}" for i, c in enumerate(r) if c) for r in sample)
        response = client.responses.create(model=model, input=_PLAN_PROMPT + table)
        out = re.sub(r"^```(?:json)?\s*|\s*```$", "", (response.output_text or "").strip())
        start, end = out.find("{"), out.rfind("}")
        doc = json.loads(out[start:end + 1]) if start >= 0 else {}
        if not isinstance(doc, dict):
            raise ValueError("no plan in the model's answer")
        return doc

    return plan
