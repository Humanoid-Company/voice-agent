"""«Знання»: a file → records → a team set in the store → Єва looks it up (server/knowledge*.py)."""
from __future__ import annotations

import asyncio
import io
import json
import os
import urllib.parse

import pytest
from fastapi.testclient import TestClient

import server.app as web
from server import knowledge, knowledge_import
from server.knowledge import KnowledgeStore, Record, _score
from server.knowledge_import import KnowledgeImportError, parse_file, table_to_items
from tools.executor import ToolExecutionContext

PEOPLE_CSV = (
    "Ім'я;Прізвище;Вік;Хобі;Улюблені фільми\n"
    "Олег;Петренко;34;риболовля, шахи;Кримінальне чтиво\n"
    "Марина;Коваль;29;йога;Амелі\n"
    ";;;;\n"
).encode()


# The same tests on Postgres (Neon's engine) when a throwaway database is given:
# KNOWLEDGE_TEST_DATABASE_URL=postgresql://… pytest tests/test_knowledge.py — its kb_* tables are dropped.
TEST_PG = os.getenv("KNOWLEDGE_TEST_DATABASE_URL", "")


@pytest.fixture(params=["sqlite", "postgres"])
def store(request, tmp_path, monkeypatch):
    if request.param == "postgres":
        if not TEST_PG:
            pytest.skip("KNOWLEDGE_TEST_DATABASE_URL not set")
        import psycopg

        with psycopg.connect(TEST_PG) as conn:
            conn.execute("DROP TABLE IF EXISTS kb_items, kb_sets")
        s = KnowledgeStore(TEST_PG)
    else:
        s = KnowledgeStore("", str(tmp_path / "kb.db"))
    monkeypatch.setattr(knowledge, "store", s)
    return s


@pytest.fixture()
def client(store, monkeypatch):
    return TestClient(web.app)


# ── files → records ──
def test_csv_semicolons_first_and_last_name_make_the_title():
    parsed = parse_file("люди.csv", PEOPLE_CSV)
    assert [i["title"] for i in parsed.items] == ["Олег Петренко", "Марина Коваль"]  # the empty row is skipped
    assert parsed.items[0]["fields"] == {"Вік": "34", "Хобі": "риболовля, шахи", "Улюблені фільми": "Кримінальне чтиво"}
    assert parsed.to_dict()["columns"] == ["Вік", "Хобі", "Улюблені фільми"]


def test_csv_in_windows_1251_and_commas():
    data = "ПІБ,Посада\nІрина Мельник,дизайнерка\n".encode("cp1251")
    assert parse_file("a.csv", data).items == [{"title": "Ірина Мельник", "fields": {"Посада": "дизайнерка"}}]


def test_no_name_column_takes_the_first():
    items = table_to_items([["Хто саме", "Хобі"], ["Тарас", "біг"]])
    assert items == [{"title": "Тарас", "fields": {"Хобі": "біг"}}]


@pytest.mark.parametrize("doc", [
    [{"name": "Олег", "age": 34, "films": ["Чтиво", "Матриця"]}],
    {"people": [{"name": "Олег", "age": 34, "films": ["Чтиво", "Матриця"]}]},
    {"Олег": {"age": 34, "films": ["Чтиво", "Матриця"]}},
])
def test_json_shapes(doc):
    items = parse_file("a.json", json.dumps(doc, ensure_ascii=False).encode()).items
    assert items == [{"title": "Олег", "fields": {"age": "34", "films": "Чтиво, Матриця"}}]


def test_xlsx():
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.append(["ПІБ", "Вік", "Хобі"])
    sheet.append(["Олег Петренко", 34, "шахи"])
    buf = io.BytesIO()
    book.save(buf)
    assert parse_file("a.xlsx", buf.getvalue()).items == [{"title": "Олег Петренко", "fields": {"Вік": "34", "Хобі": "шахи"}}]


def test_docx_table_and_text_through_the_model():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Корпоратив 20.10 починається о 19:00 у ресторані «Липа». " * 6)
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "ПІБ", "Хобі"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "Олег", "шахи"
    buf = io.BytesIO()
    doc.save(buf)
    seen = []

    def extract(text):
        seen.append(text)
        return [{"title": "Корпоратив 20.10", "fields": {"початок": "19:00", "місце": "ресторан «Липа»"}}]

    parsed = parse_file("event.docx", buf.getvalue(), extract)
    assert [i["title"] for i in parsed.items] == ["Олег", "Корпоратив 20.10"]
    assert parsed.kind == "mixed" and "Липа" in seen[0] and any("модель" in w for w in parsed.warnings)


def test_text_without_a_model_and_bad_files_say_why():
    with pytest.raises(KnowledgeImportError, match="таблицю"):
        parse_file("a.txt", "Олег любить шахи.".encode())
    with pytest.raises(KnowledgeImportError, match=r"\.docx"):
        parse_file("a.doc", b"x")
    with pytest.raises(KnowledgeImportError, match="CSV"):
        parse_file("a.pdf", b"x")
    with pytest.raises(KnowledgeImportError, match="записів"):
        parse_file("a.csv", b"only a header\n")


# ── the store ──
def test_store_create_toggle_delete(store):
    created = store.create_set("Корпоратив 20.10", parse_file("a.csv", PEOPLE_CSV).items, note="хто прийде", source="a.csv")
    assert created["count"] == 2 and created["items"][1]["title"] == "Марина Коваль"
    sets, records = store.enabled_records()
    assert sets[0]["name"] == "Корпоратив 20.10" and sets[0]["titles"] == ["Олег Петренко", "Марина Коваль"]
    assert records[0].line() == "Олег Петренко — Вік: 34; Хобі: риболовля, шахи; Улюблені фільми: Кримінальне чтиво"
    store.update_set(created["id"], enabled=False)
    assert store.enabled_records() == ([], [])  # off: Єва doesn't know it
    assert store.delete_set(created["id"]) and store.list_sets() == []


def test_a_table_from_before_owners_gets_the_columns(tmp_path):
    import sqlite3

    path = str(tmp_path / "old.db")
    with sqlite3.connect(path) as conn:  # kb_sets as prod had it before owners
        conn.execute("""CREATE TABLE kb_sets (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL,
            note TEXT NOT NULL DEFAULT '', source TEXT NOT NULL DEFAULT '', enabled INTEGER NOT NULL DEFAULT 1,
            item_count INTEGER NOT NULL DEFAULT 0, created_at REAL NOT NULL)""")
        conn.execute("INSERT INTO kb_sets (name, item_count, created_at) VALUES ('Старий', 0, 1)")
    s = KnowledgeStore("", path)
    assert s.list_sets()[0]["name"] == "Старий" and s.list_sets()[0]["owner_id"] is None
    created = s.create_set("Новий", [{"title": "Олег"}], owner_id=7, owner_name="Анна")
    assert (created["owner_id"], created["owner_name"]) == (7, "Анна")
    assert KnowledgeStore("", path).list_sets()[0]["owner_id"] == 7  # a second start: nothing added twice


def test_lookup_returns_everything_while_it_fits_else_the_closest(store, monkeypatch):
    store.create_set("Корпоратив", parse_file("a.csv", PEOPLE_CSV).items)
    found = store.lookup("Олегові")
    assert found["shown"] == found["total"] == 2 and "Марина Коваль" in found["text"]
    monkeypatch.setattr(knowledge, "LOOKUP_ALL_CHARS", 60)
    store._changed()
    found = store.lookup("що любить Марина")  # too big to send all: the closest record comes first
    assert found["shown"] == 2 and found["text"].startswith("[Корпоратив] Марина Коваль")


def test_score_matches_declined_and_misheard_names():
    oleh = Record("s", "Олег Петренко", {"хобі": "шахи"})
    maryna = Record("s", "Марина Коваль", {"хобі": "йога"})
    assert _score("Олегові", oleh) > _score("Олегові", maryna)
    assert _score("про Марину", maryna) > _score("про Марину", oleh)
    assert _score("хто грає в шахи", oleh) > _score("хто грає в шахи", maryna)


# ── HTTP: the page's two steps, the switch, delete, the access code ──
def test_upload_preview_then_save(client, store, monkeypatch):
    seen = []
    monkeypatch.setattr(web, "_plan", lambda header, sample: seen.append(header) or {})  # no model call in tests
    res = client.post("/api/knowledge/parse", content=PEOPLE_CSV,
                      headers={"X-File-Name": urllib.parse.quote("люди корпоратив.csv")})
    assert res.status_code == 200 and len(res.json()["items"]) == 2
    assert seen and seen[0][0] == "Ім'я"  # the model gets the table's headers
    assert client.get("/api/knowledge").json()["sets"] == []  # nothing saved by the preview
    items = res.json()["items"][:1]  # the user removed Марина in the preview
    saved = client.post("/api/knowledge", json={"name": "Корпоратив 20.10", "items": items, "source": "люди.csv"}).json()
    assert saved["count"] == 1
    listing = client.get("/api/knowledge").json()
    assert listing["storage"] == store.kind and listing["sets"][0]["name"] == "Корпоратив 20.10"
    assert client.post(f"/api/knowledge/{saved['id']}", json={"enabled": False}).json()["enabled"] is False
    assert client.get(f"/api/knowledge/{saved['id']}").json()["items"][0]["title"] == "Олег Петренко"
    assert client.delete(f"/api/knowledge/{saved['id']}").json() == {"ok": True}
    assert client.get(f"/api/knowledge/{saved['id']}").status_code == 404


def test_bad_file_and_empty_set_are_refused_with_a_reason(client):
    res = client.post("/api/knowledge/parse", content=b"x", headers={"X-File-Name": "a.pdf"})
    assert res.status_code == 422 and "CSV" in res.json()["detail"]
    res = client.post("/api/knowledge", json={"name": "Порожній", "items": []})
    assert res.status_code == 503 and "запису" in res.json()["detail"]


def test_knowledge_is_open_and_anyone_changes_it(client):
    assert client.get("/api/knowledge").status_code == 200
    saved = client.post("/api/knowledge", json={"name": "X", "items": [{"title": "Олег"}]})
    assert saved.status_code == 200
    set_id = saved.json()["id"]
    assert client.post(f"/api/knowledge/{set_id}", json={"enabled": False}).json()["enabled"] is False
    assert client.delete(f"/api/knowledge/{set_id}").status_code == 200
    # A set from before owners were kept goes away the same way.
    old = knowledge.store.create_set("Z", [{"title": "Тарас"}])
    assert client.delete(f"/api/knowledge/{old['id']}").status_code == 200
    # A new event: everyone may come in again.
    assert client.delete("/api/concierge/checkins").status_code == 200
    assert client.get("/api/concierge/checkins").status_code == 200


# ── Єва ──
def test_the_tool_and_the_prompts(store):
    user = web.users.get("kb-browser-0123456789abc")
    ctx = ToolExecutionContext(session_id="s")
    empty = asyncio.run(user.executor.execute("knowledge_lookup", {"query": "Олег"}, ctx))
    assert empty.status == "empty" and "Знання" in empty.message
    store.create_set("Корпоратив 20.10", parse_file("a.csv", PEOPLE_CSV).items)
    found = asyncio.run(user.executor.execute("knowledge_lookup", {"query": "Олегові", "set": "корпоратив"}, ctx))
    assert found.ok and "Олег Петренко — Вік: 34" in found.message and "лише з них" in found.message
    # The assistant knows people like an acquaintance: lookup + remember_person, a note on how, no check-in.
    plain = web._session_config(user)["delegation"]["responses"]
    names = {t["name"] for t in plain["tools"]}
    assert {"knowledge_lookup", "remember_person"} <= names and not names & {"find_guest", "confirm_guest"}
    assert "PEOPLE YOU KNOW" in plain["instructions"] and "no check question" in plain["instructions"]
    backend = web._session_config(user, concierge=True)["delegation"]["responses"]
    assert "knowledge_lookup" in backend["instructions"] and any(t["name"] == "knowledge_lookup" for t in backend["tools"])
    assert "remember_person" not in {t["name"] for t in backend["tools"]} and "PEOPLE YOU KNOW" not in backend["instructions"]


def test_eva_writes_down_people_and_adds_to_them(store):
    user = web.users.get("kb-browser-0123456789abc")
    ctx = ToolExecutionContext(session_id="s")

    def run(tool, args):
        return asyncio.run(user.executor.execute(tool, args, ctx))

    first = run("remember_person", {"name": "Іра", "facts": {"Хто це": "сестра користувача", "Вік": "30"}})
    assert first.ok and first.data["new"] is True
    again = run("remember_person", {"name": "іра", "facts": {"Хобі": "йога", "Вік": "30"}})
    assert again.ok and again.data["new"] is False  # same person: facts added, not a second record
    more = run("remember_person", {"name": "Ірина Коваль", "facts": {"Хобі": "біг"}})
    assert more.data["new"] is True  # another name is another person, never a guessed merge
    sets = store.list_sets()
    assert [s["name"] for s in sets] == [knowledge.MEMORY_SET]
    assert sets[0]["count"] == 2 == len(store.get_set(sets[0]["id"])["items"])
    found = run("knowledge_lookup", {"query": "Іра"})
    assert "сестра користувача" in found.message and "Хобі: йога" in found.message
    # Nothing is lost: a second value under the same label goes after the first.
    run("remember_person", {"name": "Іра", "facts": {"Хобі": "малювання"}})
    assert "йога; малювання" in run("knowledge_lookup", {"query": "Іра"}).message
    assert not run("remember_person", {"name": "Іра", "facts": {}}).ok  # nothing to write


def test_what_eva_wrote_down_is_never_the_guest_list(store):
    user = web.users.get("kb-browser-0123456789abc")
    ctx = ToolExecutionContext(session_id="s")
    asyncio.run(user.executor.execute("remember_person", {"name": "Іван Шпак", "facts": {"Хобі": "шахи"}}, ctx))
    store.create_set("Корпоратив", parse_file("a.csv", PEOPLE_CSV).items)
    user.concierge = True
    found = asyncio.run(user.executor.execute("find_guest", {"name": "Іван", "surname": "Шпак"}, ctx))
    assert found.status == "not_found"  # written down in a talk is not on the list
    confirmed = asyncio.run(user.executor.execute(
        "confirm_guest", {"full_name": "Іван Шпак", "verified": True, "guest_answer": "шахи"}, ctx))
    assert confirmed.status == "not_on_list"
    assert "Шпак" not in asyncio.run(user.executor.execute("knowledge_lookup", {"query": "Іван"}, ctx)).message
    assert "Петренко" in asyncio.run(user.executor.execute("knowledge_lookup", {"query": "Олег"}, ctx)).message


def test_fast_lookup_ranks_like_scoring_record_by_record():
    """_closest compares each word pair once for the whole list (2000 people: 1.6 s → 11 ms a lookup); it must
    rank exactly as scoring every record on its own."""
    people = [("Олег Петренко", "шахи"), ("Олег Гнатюк", "футбол"), ("Марина Коваль", "йога"), ("Тарас Бойко", "біг"),
              ("Ірина Мельник", "настолки"), ("Катерина Шаповал", "книжки")] * 5
    records = [Record("s", t, {"Хобі": h, "Посада": "розробник"}) for t, h in people]
    from server.knowledge import _closest

    for query in ("Олегові шахи", "хто любить йогу", "Алєг", "Катерину", "розробник", ""):
        slow = sorted(records, key=lambda r: -_score(query, r))[:7]
        assert [r.line() for r in _closest(query, records, 7)] == [r.line() for r in slow]


def test_a_form_export_titles_records_by_the_name_not_the_timestamp():
    rows = [["Позначка часу", "Твоє Ім'я та прізвище", "Улюблений фільм"],
            ["2026-10-19 13:51", "Богдан Рудяк", "Сімпсони"], ["2026-10-19 13:56", "Олександра Іваченко", "Декстер"]]
    assert [i["title"] for i in knowledge_import.table_to_items(rows)] == ["Богдан Рудяк", "Олександра Іваченко"]
    # No name column at all: the first column that is not a time.
    plain = [["Timestamp", "Нік", "Хобі"], ["2026-10-19", "Вовк", "шахи"]]
    assert knowledge_import.table_to_items(plain)[0]["title"] == "Вовк"


FORM_ROWS = [
    ["Позначка часу", "Твоє Ім'я та прізвище", "Який твій улюблений фільм, серіал або гра?", "Два цікаві факти про себе"],
    ["2026-10-19 13:51", "Богдан Рудяк", "Сімпсони", "Люблю спати"],
]


def test_the_model_names_the_title_and_shortens_question_headers():
    asked = []

    def plan(header, sample):
        asked.append((header, sample))
        return {"title": [1], "skip": [0],
                "labels": {"2": "Улюблений фільм, серіал, гра?", "3": "Цікаві факти", "1": "x"}}

    [item] = knowledge_import.table_to_items(FORM_ROWS, plan)
    assert item == {"title": "Богдан Рудяк", "fields": {"Улюблений фільм, серіал, гра": "Сімпсони", "Цікаві факти": "Люблю спати"}}
    assert asked[0][0][2].startswith("Який твій") and asked[0][1] == [FORM_ROWS[1]]


def test_a_bad_plan_falls_back_to_the_plain_rules():
    def broken(header, sample):
        raise RuntimeError("model down")

    wrong = lambda header, sample: {"title": [99, True, "1"], "labels": {"7": "x", "2": "  "}, "skip": "all"}  # noqa: E731
    for plan in (broken, wrong, lambda h, s: "nonsense"):
        [item] = knowledge_import.table_to_items(FORM_ROWS, plan)
        assert item["title"] == "Богдан Рудяк"  # the loose name rule still finds it
        assert "Який твій улюблений фільм, серіал або гра?" in item["fields"]
    # Labels never collide: two columns shortened to the same words stay two facts.
    twin = lambda h, s: {"labels": {"2": "Фільм", "3": "Фільм"}}  # noqa: E731
    assert list(knowledge_import.table_to_items(FORM_ROWS, twin)[0]["fields"]) == ["Позначка часу", "Фільм", "Фільм (2)"]


# ── text sections and «Завжди пам'ятати» ──
RULES = "## Як говорити\nКоротко, 1–3 речення.\n\n## Теми\nНе говори про політику."
STORIES = "Вступ без заголовка.\n\n## Кавуни\nДавид виростив кавуни з бабусею.\n\n## Коти\nЮліко годує котів."


def test_text_sections_split_at_headings():
    records = knowledge.text_sections("Історії", STORIES)
    assert [r.title for r in records] == ["Історії", "Кавуни", "Коти"]
    assert records[1].fields == {"Текст": "Давид виростив кавуни з бабусею."}
    assert [r.title for r in knowledge.text_sections("Нотатка", "просто текст")] == ["Нотатка"]


def test_a_text_set_is_looked_up_by_section_and_is_never_a_guest(store):
    store.create_set("Гості", parse_file("a.csv", PEOPLE_CSV).items)
    created = store.create_set("Історії", [], kind="text", body=STORIES + "\n\n\n\n  ")
    assert created["kind"] == "text" and created["count"] == 0 and created["body"].endswith("котів.")
    assert "body" not in store.list_sets()[0] and store.list_sets()[0]["chars"] == len(created["body"])
    found = store.lookup("кавуни")
    assert "[Історії] Кавуни — Текст: Давид виростив кавуни" in found["text"] and "Олег Петренко" in found["text"]
    sets, people = store.enabled_records()  # the concierge's guest list: people only
    assert [s["name"] for s in sets] == ["Гості"] and {r.title for r in people} == {"Олег Петренко", "Марина Коваль"}
    store.update_set(created["id"], body="## Нове\nІнша історія.")
    assert "Інша історія" in store.lookup("історія")["text"] and "кавуни" not in store.lookup("кавуни")["text"]
    with pytest.raises(knowledge.KnowledgeError, match="порожній"):
        store.update_set(created["id"], body="   ")
    with pytest.raises(knowledge.KnowledgeError, match="порожній"):
        store.create_set("Порожній", [], kind="text", body="")
    people_set = store.list_sets()[-1]
    with pytest.raises(knowledge.KnowledgeError, match="текстовому"):
        store.update_set(people_set["id"], body="текст")


def test_pinned_sets_go_to_the_prompt_not_to_the_lookup(store, monkeypatch):
    rules = store.create_set("Правила свята", [], kind="text", body=RULES, pinned=True)
    guests = store.create_set("Гості", parse_file("a.csv", PEOPLE_CSV).items)
    assert store.pinned_block() == "## Правила свята\n" + RULES.replace("## ", "### ")
    assert "Правила" not in store.lookup("політика")["text"]  # in the instructions already
    store.update_set(guests["id"], pinned=True)  # a short guest table, pinned too
    block = store.pinned_block()
    assert block.startswith("## Правила свята") and "## Гості\nОлег Петренко — Вік: 34" in block
    assert store.lookup("Олег")["total"] == 0
    assert len(store.enabled_records()[1]) == 2  # still the guest list
    store.update_set(rules["id"], enabled=False)  # off: not in the prompt either
    assert "Правила" not in store.pinned_block()
    monkeypatch.setattr(knowledge, "PINNED_MAX_CHARS", 30)
    store._changed()
    assert store.pinned_block().endswith("забагато в «Завжди пам'ятати»]")


def test_evas_own_set_is_never_pinned_or_renamed(store):
    store.remember("Іра", {"Хобі": "йога"})
    own = store.list_sets()[0]
    with pytest.raises(knowledge.KnowledgeError):
        store.update_set(own["id"], pinned=True)
    with pytest.raises(knowledge.KnowledgeError):
        store.update_set(own["id"], name="Інше")
    with pytest.raises(knowledge.KnowledgeError):
        store.create_set(knowledge.MEMORY_SET, [{"title": "Олег"}])


def test_a_database_from_before_text_sets_gets_the_columns(tmp_path):
    import sqlite3

    path = str(tmp_path / "old.db")
    with sqlite3.connect(path) as conn:  # kb_sets as prod has it now
        conn.execute("""CREATE TABLE kb_sets (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL,
            note TEXT NOT NULL DEFAULT '', source TEXT NOT NULL DEFAULT '', enabled INTEGER NOT NULL DEFAULT 1,
            item_count INTEGER NOT NULL DEFAULT 0, created_at REAL NOT NULL, owner_id INTEGER,
            owner_name TEXT NOT NULL DEFAULT '')""")
        conn.execute("INSERT INTO kb_sets (name, item_count, created_at) VALUES ('Старий', 0, 1)")
    s = KnowledgeStore("", path)
    old = s.list_sets()[0]
    assert (old["kind"], old["pinned"], old["chars"]) == ("records", False, 0)
    s.create_set("Правила", [], kind="text", body="Коротко.", pinned=True)
    assert s.pinned_block() == "## Правила\nКоротко."


def test_http_text_sections_and_pinning(client, store):
    saved = client.post("/api/knowledge", json={"name": "Правила свята", "kind": "text", "body": RULES, "pinned": True})
    assert saved.status_code == 200 and saved.json()["pinned"] is True
    listing = client.get("/api/knowledge").json()
    assert listing["pinned_chars"] == len("## Правила свята\n" + RULES.replace("## ", "### ")) and listing["pinned_max"] == knowledge.PINNED_MAX_CHARS
    set_id = saved.json()["id"]
    assert client.get(f"/api/knowledge/{set_id}").json()["body"] == RULES
    changed = client.post(f"/api/knowledge/{set_id}", json={"body": "Нові правила.", "name": "Правила"}).json()
    assert changed["body"] == "Нові правила." and changed["name"] == "Правила"
    assert client.post(f"/api/knowledge/{set_id}", json={"pinned": False}).json()["pinned"] is False
    assert client.get("/api/knowledge").json()["pinned_chars"] == 0
    assert client.post("/api/knowledge", json={"name": "X", "kind": "text", "body": ""}).status_code == 503
    assert client.post("/api/knowledge", json={"name": "X", "kind": "щось"}).status_code == 503
    assert client.post("/api/knowledge/999", json={"pinned": True}).status_code == 404


def test_pinned_knowledge_is_in_every_call(store):
    user = web.users.get("kb-browser-0123456789abc")
    plain = web._session_config(user)
    assert "STANDING KNOWLEDGE" not in plain["instructions"]  # nothing pinned: the prompt as before
    store.create_set("Правила свята", [], kind="text", body=RULES, pinned=True)
    plain = web._session_config(user)
    assert plain["instructions"].rstrip().endswith("Не говори про політику.")  # last: over the general rules
    assert "Не говори про політику." in plain["delegation"]["responses"]["instructions"]
    door = web._session_config(user, concierge=True)
    assert "Не говори про політику." in door["instructions"]
    assert "Не говори про політику." in door["delegation"]["responses"]["instructions"]


def test_a_store_that_is_down_never_breaks_a_call(store, monkeypatch):
    def down():
        raise knowledge.KnowledgeError("База знань недоступна")

    monkeypatch.setattr(store, "pinned_block", down)
    assert knowledge.pinned_note() == ""
    assert web._session_config(web.users.get("kb-browser-0123456789abc"))["instructions"]


def test_a_knowledge_password_closes_the_page_but_not_eva(client, store, monkeypatch):
    monkeypatch.setattr(web, "KNOWLEDGE_TOKEN", "sekret-123")
    assert client.get("/api/knowledge").status_code == 401
    assert client.post("/api/knowledge", json={"name": "X", "kind": "text", "body": "т"}).status_code == 401
    assert client.post("/api/knowledge/parse", content=PEOPLE_CSV, headers={"X-File-Name": "a.csv"}).status_code == 401
    good = {"X-Knowledge-Token": "sekret-123"}
    saved = client.post("/api/knowledge", json={"name": "X", "kind": "text", "body": "т"}, headers=good)
    assert saved.status_code == 200
    set_id = saved.json()["id"]
    for call in (client.get, client.delete):
        assert call(f"/api/knowledge/{set_id}", headers={"X-Knowledge-Token": "wrong"}).status_code == 401
    assert client.post(f"/api/knowledge/{set_id}", json={"pinned": True}).status_code == 401
    assert client.get("/api/knowledge", headers=good).json()["sets"][0]["name"] == "X"
    # Єва's own lookup goes past the password: it runs on the server.
    user = web.users.get("kb-browser-0123456789abc")
    found = asyncio.run(user.executor.execute("knowledge_lookup", {"query": "т"}, ToolExecutionContext(session_id="s")))
    assert found.ok
